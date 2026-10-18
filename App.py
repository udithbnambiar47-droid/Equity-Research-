"""
Equity Research App (single-file version): DCF, scenarios, Monte Carlo, relative valuation, export.
Run: streamlit run app.py   |   Educational use only. Not investment advice.
"""
from __future__ import annotations



from dataclasses import dataclass, replace

import numpy as np

TV_METHODS = ("gordon", "exit", "blend")


@dataclass(frozen=True)
class Assumptions:
    rev0: float              # base-year revenue (or base FCF in fallback mode)
    growth_start: float      # year-1 revenue growth; fades linearly to terminal_g
    margin_start: float      # current EBIT margin
    margin_target: float     # EBIT margin reached after `margin_years`
    margin_years: int
    tax: float               # tax rate on EBIT
    da_pct: float            # D&A as % of revenue
    capex_pct: float         # capex as % of revenue
    nwc_pct: float           # investment in NWC as % of incremental revenue
    wacc: float
    terminal_g: float
    exit_multiple: float     # EV / EBITDA applied to final-year EBITDA
    years: int
    net_debt: float
    shares: float
    tv_method: str = "gordon"  # gordon | exit | blend


def _c(x):
    """Parameter -> column vector (n, 1)."""
    return np.atleast_1d(np.asarray(x, dtype=float))[:, None]


def run_dcf(a: Assumptions) -> dict:
    """Run the DCF. Returns 2-D arrays (n, years) and 1-D arrays (n,)."""
    T = int(a.years)
    t = np.arange(1, T + 1, dtype=float)[None, :]

    g0, tg, w = _c(a.growth_start), _c(a.terminal_g), _c(a.wacc)
    m0, m1 = _c(a.margin_start), _c(a.margin_target)
    tax, da_p, cx_p, nwc, mult = (_c(a.tax), _c(a.da_pct), _c(a.capex_pct),
                                  _c(a.nwc_pct), _c(a.exit_multiple))

    growth = g0 + (tg - g0) * (t - 1) / max(T - 1, 1)          # fades to terminal g
    k = max(min(int(a.margin_years), T), 1)
    margin = m0 + (m1 - m0) * np.minimum(t / k, 1.0)           # linear ramp to target

    rev = float(a.rev0) * np.cumprod(1.0 + growth, axis=1)
    prev = np.concatenate([np.full((rev.shape[0], 1), float(a.rev0)), rev[:, :-1]], axis=1)

    ebit = rev * margin
    da = rev * da_p
    capex = rev * cx_p
    d_nwc = nwc * (rev - prev)
    fcff = ebit * (1.0 - tax) + da - capex - d_nwc
    ebitda = ebit + da

    disc = (1.0 + w) ** t
    pv = fcff / disc

    f_T, e_T = fcff[:, -1:], ebitda[:, -1:]
    valid = w > tg
    den = np.where(valid, w - tg, np.nan)
    with np.errstate(divide="ignore", invalid="ignore"):
        tv_g = f_T * (1.0 + tg) / den                           # Gordon growth
        tv_x = e_T * mult                                       # exit multiple
        implied_g = (tv_x * w - f_T) / (tv_x + f_T)             # g implied by exit TV
        implied_mult = tv_g / e_T                               # multiple implied by Gordon TV

    if a.tv_method == "gordon":
        tv = tv_g
    elif a.tv_method == "exit":
        tv = tv_x
    else:
        tv = (tv_g + tv_x) / 2.0

    pv_tv = tv / disc[:, -1:]
    ev = pv.sum(axis=1, keepdims=True) + pv_tv
    equity = ev - a.net_debt
    per_share = equity / a.shares

    one = lambda x: x[:, 0]  # noqa: E731
    with np.errstate(divide="ignore", invalid="ignore"):
        tv_share = pv_tv / ev
    return {
        "years": T,
        "growth": growth, "margin": margin, "rev": rev, "ebit": ebit, "da": da,
        "capex": capex, "d_nwc": d_nwc, "fcff": fcff, "ebitda": ebitda,
        "disc": disc, "pv": pv,
        "tv_gordon": one(tv_g), "tv_exit": one(tv_x), "tv": one(tv), "pv_tv": one(pv_tv),
        "ev": one(ev), "equity": one(equity), "per_share": one(per_share),
        "implied_g": one(implied_g), "implied_mult": one(implied_mult),
        "tv_share": one(tv_share), "pv_sum": pv.sum(axis=1),
    }


def point_value(a: Assumptions) -> float:
    """Intrinsic value per share for scalar assumptions."""
    return float(run_dcf(a)["per_share"][0])


# --------------------------------------------------------------------------
# Scenarios
# --------------------------------------------------------------------------
def apply_scenario(a: Assumptions, d_growth=0.0, d_margin=0.0, d_wacc=0.0,
                   d_tg=0.0, d_mult=0.0) -> Assumptions:
    """Shift the key drivers (all deltas in decimals, multiple delta in turns)."""
    return replace(
        a,
        growth_start=a.growth_start + d_growth,
        margin_target=a.margin_target + d_margin,
        wacc=a.wacc + d_wacc,
        terminal_g=max(a.terminal_g + d_tg, 0.0),
        exit_multiple=max(a.exit_multiple + d_mult, 1.0),
    )


# --------------------------------------------------------------------------
# Monte Carlo
# --------------------------------------------------------------------------
def monte_carlo(a: Assumptions, n: int, seed: int, sd: dict) -> np.ndarray:
    """
    Simulate intrinsic value per share. Each key input is drawn independently
    from a normal distribution centred on the base case, then clipped to
    sensible bounds. `sd` keys: growth, margin, wacc, tg (decimals), mult_rel (fraction).
    """
    rng = np.random.default_rng(seed)
    g = np.clip(rng.normal(a.growth_start, sd["growth"], n), -0.20, 0.60)
    m = np.clip(rng.normal(a.margin_target, sd["margin"], n), -0.10, max(0.70, a.margin_target))
    w = np.clip(rng.normal(a.wacc, sd["wacc"], n), 0.04, 0.25)
    tg = np.clip(rng.normal(a.terminal_g, sd["tg"], n), 0.0, 0.05)
    tg = np.minimum(tg, w - 0.01)                       # keep WACC > g
    mult = np.maximum(rng.normal(a.exit_multiple, a.exit_multiple * sd["mult_rel"], n), 1.0)
    res = run_dcf(replace(a, growth_start=g, margin_target=m, wacc=w,
                          terminal_g=tg, exit_multiple=mult))
    out = res["per_share"]
    return out[np.isfinite(out)]


def mc_summary(samples: np.ndarray, price: float, band: float) -> dict:
    pct = {p: float(np.percentile(samples, p)) for p in (5, 10, 25, 50, 75, 90, 95)}
    return {
        "mean": float(samples.mean()), "std": float(samples.std()), "pct": pct,
        "p_above": float((samples > price).mean()),
        "p_above_band": float((samples > price * (1 + band)).mean()),
        "p_below_band": float((samples < price * (1 - band)).mean()),
        "n": int(samples.size),
    }


# --------------------------------------------------------------------------
# Sensitivity
# --------------------------------------------------------------------------
def sensitivity(a: Assumptions, axis: str = "tg", n: int = 5):
    """Grid of value/share: WACC (rows) vs terminal growth or exit multiple (cols)."""
    steps = np.linspace(-(n // 2), n // 2, n)
    waccs = a.wacc + steps * 0.01
    if axis == "tg":
        cols = np.maximum(a.terminal_g + steps * 0.005, 0.0)
    else:
        cols = np.maximum(a.exit_multiple + steps * 1.0, 1.0)
    grid = np.full((n, n), np.nan)
    for i, w in enumerate(waccs):
        for j, c in enumerate(cols):
            kw = {"terminal_g": float(c)} if axis == "tg" else {"exit_multiple": float(c)}
            grid[i, j] = point_value(replace(a, wacc=float(w), **kw))
    return waccs, cols, grid


# --------------------------------------------------------------------------
# WACC and formatting helpers
# --------------------------------------------------------------------------
def compute_wacc(rf, beta, erp, cost_debt, tax, mcap, debt):
    ke = rf + beta * erp
    kd = cost_debt * (1 - tax)
    total = mcap + debt
    we, wd = (mcap / total, debt / total) if total > 0 else (1.0, 0.0)
    return we * ke + wd * kd, ke, kd, we, wd


def fmt_big(x, cur=""):
    if x is None or not np.isfinite(x):
        return "n/a"
    for unit, div in (("T", 1e12), ("B", 1e9), ("M", 1e6)):
        if abs(x) >= div:
            return f"{cur}{x / div:,.2f}{unit}"
    return f"{cur}{x:,.0f}"




from dataclasses import dataclass, field

import numpy as np
import pandas as pd
import streamlit as st
import yfinance as yf

# A few starter peer sets. Anything else: the user types their own peers.
DEFAULT_PEERS = {
    "AAPL": ["MSFT", "GOOGL", "META", "AMZN"],
    "MSFT": ["AAPL", "GOOGL", "ORCL", "ADBE"],
    "TCS.NS": ["INFY.NS", "HCLTECH.NS", "WIPRO.NS", "TECHM.NS"],
    "INFY.NS": ["TCS.NS", "HCLTECH.NS", "WIPRO.NS", "TECHM.NS"],
    "HDFCBANK.NS": ["ICICIBANK.NS", "KOTAKBANK.NS", "AXISBANK.NS", "SBIN.NS"],
    "ICICIBANK.NS": ["HDFCBANK.NS", "KOTAKBANK.NS", "AXISBANK.NS", "SBIN.NS"],
}


def _series(df, *names) -> pd.Series:
    """First matching row of a yfinance statement as a year-indexed series (oldest first)."""
    if df is None or getattr(df, "empty", True):
        return pd.Series(dtype=float)
    for n in names:
        if n in df.index:
            s = pd.to_numeric(df.loc[n], errors="coerce").dropna().astype(float)
            s.index = pd.to_datetime(s.index).year
            s = s[~s.index.duplicated(keep="first")]
            return s.sort_index()
    return pd.Series(dtype=float)


@st.cache_data(ttl=3600, show_spinner=False)
def load_company(ticker: str) -> dict:
    t = yf.Ticker(ticker)
    return {
        "info": t.info or {},
        "cashflow": t.cashflow,
        "income": t.income_stmt,
    }


@st.cache_data(ttl=3600, show_spinner=False)
def load_multiples(tickers: tuple) -> pd.DataFrame:
    rows = []
    for tk in tickers:
        try:
            i = yf.Ticker(tk).info or {}
        except Exception:
            i = {}
        rows.append({
            "Ticker": tk,
            "Name": i.get("shortName") or tk,
            "Market cap": i.get("marketCap"),
            "P/E (TTM)": i.get("trailingPE"),
            "P/E (Fwd)": i.get("forwardPE"),
            "EV/EBITDA": i.get("enterpriseToEbitda"),
            "EV/Revenue": i.get("enterpriseToRevenue"),
        })
    df = pd.DataFrame(rows)
    for c in ("Market cap", "P/E (TTM)", "P/E (Fwd)", "EV/EBITDA", "EV/Revenue"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df.replace([np.inf, -np.inf], np.nan)


@dataclass
class Defaults:
    has_operating: bool
    rev0: float = np.nan
    growth: float = 0.08
    margin_now: float = np.nan
    margin_target: float = np.nan
    da_pct: float = 0.04
    capex_pct: float = 0.05
    nwc_pct: float = 0.05
    tax: float = 0.21
    ebitda_ttm: float = np.nan
    eps: float = np.nan
    fcf_hist: pd.Series = field(default_factory=lambda: pd.Series(dtype=float))
    rev_hist: pd.Series = field(default_factory=lambda: pd.Series(dtype=float))
    margin_hist: pd.Series = field(default_factory=lambda: pd.Series(dtype=float))
    fcf_year: int = 0


def _tail_mean(s: pd.Series, n: int = 3, default: float = np.nan) -> float:
    s = s.replace([np.inf, -np.inf], np.nan).dropna()
    return float(s.tail(n).mean()) if len(s) else default


def _fcf_history(cf) -> pd.Series:
    s = _series(cf, "Free Cash Flow")
    if s.empty:
        ocf = _series(cf, "Operating Cash Flow", "Total Cash From Operating Activities")
        capex = _series(cf, "Capital Expenditure")
        if ocf.empty or capex.empty:
            return pd.Series(dtype=float)
        s = (ocf + capex).dropna()  # capex is negative in yfinance
    return s


def derive(company: dict) -> Defaults:
    """Derive sensible starting assumptions from reported history."""
    info, cf, inc = company["info"], company["cashflow"], company["income"]

    rev = _series(inc, "Total Revenue", "Operating Revenue")
    ebit = _series(inc, "EBIT", "Operating Income")
    da = _series(cf, "Depreciation And Amortization", "Depreciation Amortization Depletion", "Depreciation")
    capex = _series(cf, "Capital Expenditure")
    d_wc = _series(cf, "Change In Working Capital")
    tax_rate = _series(inc, "Tax Rate For Calcs")
    fcf = _fcf_history(cf)

    margin = (ebit / rev).dropna() if len(rev) and len(ebit) else pd.Series(dtype=float)
    has_op = bool(len(rev) and len(margin) and rev.iloc[-1] > 0)

    d = Defaults(has_operating=has_op, fcf_hist=fcf, rev_hist=rev, margin_hist=margin)
    d.fcf_year = int(fcf.index[-1]) if len(fcf) else 0
    d.eps = info.get("trailingEps", np.nan) or np.nan

    if has_op:
        d.rev0 = float(rev.iloc[-1])
        pos = rev[rev > 0]
        if len(pos) >= 2 and pos.index[-1] > pos.index[0]:
            n = pos.index[-1] - pos.index[0]
            d.growth = float(np.clip((pos.iloc[-1] / pos.iloc[0]) ** (1 / n) - 1, -0.05, 0.30))
        d.margin_now = float(margin.iloc[-1])
        d.margin_target = float(np.clip(_tail_mean(margin, 3, d.margin_now), -0.10, 0.60))
        d.da_pct = float(np.clip(_tail_mean(da / rev, 3, 0.04), 0.0, 0.30))
        d.capex_pct = float(np.clip(_tail_mean(capex.abs() / rev, 3, 0.05), 0.0, 0.40))
        d.tax = float(np.clip(_tail_mean(tax_rate, 3, 0.21), 0.05, 0.40))
        nwc = (-d_wc / rev.diff()).replace([np.inf, -np.inf], np.nan).dropna()
        d.nwc_pct = float(np.clip(nwc.median(), 0.0, 0.30)) if len(nwc) else 0.05

    ebitda = info.get("ebitda")
    if not ebitda and has_op and len(da):
        ebitda = float(ebit.iloc[-1] + da.iloc[-1])
    d.ebitda_ttm = float(ebitda) if ebitda else np.nan
    return d




import io
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402



XL_FORMATS = {"big": "#,##0", "price": "#,##0.00", "num": "#,##0.00",
              "pct": "0.0%", "x": '0.0"x"', "int": "0", "txt": "@"}


def _fmt_val(v, kind: str, cur: str = "") -> str:
    if v is None or (isinstance(v, (float, np.floating)) and not np.isfinite(v)):
        return "n/a"
    if pd.isna(v) if not isinstance(v, str) else False:
        return "n/a"
    try:
        if kind == "big":
            return fmt_big(float(v), cur)
        if kind == "price":
            return f"{cur}{float(v):,.2f}"
        if kind == "num":
            return f"{float(v):,.2f}"
        if kind == "pct":
            return f"{float(v):.1%}"
        if kind == "x":
            return f"{float(v):.1f}x"
        if kind == "int":
            return f"{int(v)}"
    except (TypeError, ValueError):
        pass
    return str(v)


def format_df(df: pd.DataFrame, spec: dict, cur: str = "") -> pd.DataFrame:
    """Return a copy of df as display strings according to spec."""
    out = pd.DataFrame(index=df.index)
    for c in df.columns:
        kind = spec.get(c, "txt")
        out[c] = [_fmt_val(v, kind, cur) for v in df[c]]
    return out


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def build_excel(report: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")

    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        summary = pd.DataFrame(report["summary"], columns=["Item", "Value"])
        summary.to_excel(xw, sheet_name="Summary", index=False)
        used = {"Summary"}
        specs = {"Summary": {}}

        for t in report["tables"]:
            name = t["name"][:31]
            while name in used:
                name = name[:28] + "_2"
            used.add(name)
            specs[name] = t["spec"]
            t["df"].to_excel(xw, sheet_name=name, index=False)

        for name, ws in xw.sheets.items():
            spec = specs.get(name, {})
            for cell in ws[1]:
                cell.fill, cell.font = head_fill, head_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            headers = [c.value for c in ws[1]]
            for j, h in enumerate(headers, start=1):
                kind = spec.get(h, "txt")
                width = len(str(h))
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell = row[0]
                    if kind != "txt" and isinstance(cell.value, (int, float)):
                        cell.number_format = XL_FORMATS[kind]
                    width = max(width, len(str(cell.value)) if cell.value is not None else 0)
                ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = min(max(width + 3, 10), 60)
    return bio.getvalue()


# --------------------------------------------------------------------------
# Charts for the PDF
# --------------------------------------------------------------------------
def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def football_png(rows, price, cur):
    fig, ax = plt.subplots(figsize=(7.2, 0.55 * len(rows) + 1.0))
    for i, (label, lo, mid, hi) in enumerate(rows):
        ax.barh(i, hi - lo, left=lo, color="#3B6FB6", height=0.5)
        if mid is not None and np.isfinite(mid):
            ax.plot([mid, mid], [i - 0.25, i + 0.25], color="white", lw=2)
        ax.text(hi, i, f"  {lo:,.0f} - {hi:,.0f}", va="center", fontsize=7)
    ax.axvline(price, color="#C0392B", ls="--", lw=1.2)
    ax.text(price, -0.8, f"Price {price:,.2f}", color="#C0392B", fontsize=7, ha="center")
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels([r[0] for r in rows], fontsize=8)
    ax.invert_yaxis()
    ax.set_xlabel(f"Value per share ({cur.strip()})", fontsize=8)
    ax.tick_params(axis="x", labelsize=7)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _png(fig)


def mc_png(samples, price, cur):
    lo, hi = np.percentile(samples, [1, 99])
    fig, ax = plt.subplots(figsize=(7.2, 2.6))
    ax.hist(np.clip(samples, lo, hi), bins=60, color="#3B6FB6")
    ax.axvline(price, color="#C0392B", ls="--", lw=1.2, label=f"Price {price:,.2f}")
    ax.axvline(np.median(samples), color="black", lw=1.2, label=f"Median {np.median(samples):,.2f}")
    ax.legend(fontsize=7, frameon=False)
    ax.set_xlabel(f"Intrinsic value per share ({cur.strip()})", fontsize=8)
    ax.tick_params(labelsize=7)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _png(fig)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def build_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    cur = report["currency"] + " "          # Helvetica has no currency glyphs, use the code
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=17, spaceAfter=2, alignment=0)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceBefore=10, spaceAfter=4,
                        textColor=colors.HexColor("#1F3A5F"))
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7.5, textColor=colors.grey)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9)

    def safe(x):
        """Helvetica only covers Latin-1; swap currency symbols and unknown glyphs."""
        x = str(x).replace("₹", "INR ").replace("€", "EUR ").replace("Δ", "Chg ")
        return x.encode("latin-1", "replace").decode("latin-1")

    def table(df, spec, col_align_left=1):
        fdf = format_df(df, spec, cur)
        data = [[safe(c) for c in fdf.columns]] + [[safe(v) for v in row] for row in fdf.values.tolist()]
        tb = Table(data, repeatRows=1, hAlign="LEFT")
        st_ = [
            ("FONT", (0, 0), (-1, -1), "Helvetica", 7.5),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (col_align_left, 0), (-1, -1), "RIGHT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
            ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.grey),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]
        tb.setStyle(TableStyle(st_))
        return tb

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=report["title"], author="DCF Equity Research App")
    story = [
        Paragraph(escape(safe(report["title"])), h1),
        Paragraph(escape(safe(f"{report['name']} ({report['ticker']})  |  {report['date']}  |  prices in {report['currency']}")), small),
        Spacer(1, 6),
        Paragraph(f"<b>Verdict: {escape(report['verdict'])}</b>", ParagraphStyle(
            "v", parent=body, fontSize=13, textColor=colors.HexColor(report.get("verdict_color", "#1F3A5F")))),
        Spacer(1, 4),
    ]

    sm = pd.DataFrame(report["summary"], columns=["Item", "Value"])
    story.append(table(sm, {}))

    if report.get("football"):
        story += [Paragraph("Valuation summary (football field)", h2),
                  Image(football_png(report["football"], report["price"], cur), width=175 * mm,
                        height=(0.55 * len(report["football"]) + 1.0) / 7.2 * 175 * mm * 0.97)]
    samples = report.get("mc_samples")
    if samples is not None and len(samples):
        story += [Paragraph("Monte Carlo distribution", h2),
                  Image(mc_png(samples, report["price"], cur), width=175 * mm, height=175 * mm * 2.6 / 7.2)]

    for t in report["tables"]:
        if not t.get("pdf", True):
            continue
        story.append(KeepTogether([Paragraph(escape(safe(t["name"])), h2), table(t["df"], t["spec"])]))

    story += [Spacer(1, 10), Paragraph(
        "Educational use only; not investment advice. Data from Yahoo Finance (yfinance) may be delayed "
        "or incomplete. Model output depends entirely on the stated assumptions.", small)]
    doc.build(story)
    return buf.getvalue()




import re
from dataclasses import replace
from datetime import date

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st





st.set_page_config(page_title="DCF Equity Research", page_icon="📈", layout="wide")

NAVY, BLUE, RED, GREY = "#1F3A5F", "#3B6FB6", "#C0392B", "#8A94A6"

st.title("📈 Equity Research: DCF, Scenarios & Relative Valuation")
st.caption("Revenue and margin driven DCF with exit-multiple cross-check, bull/base/bear scenarios, "
           "Monte Carlo and peer multiples. Educational use only, not investment advice.")

# ==========================================================================
# Sidebar: stock, forecast, discount rate, simulation, peers
# ==========================================================================
with st.sidebar:
    st.header("1. Stock")
    ticker = st.text_input("Ticker symbol", value="AAPL",
                           help="US: AAPL | NSE: TCS.NS | BSE: TCS.BO | London: VOD.L").strip().upper()
    if not ticker:
        st.stop()
    K = lambda s: f"{ticker}:{s}"  # noqa: E731  widget keys reset when the ticker changes

    with st.spinner(f"Fetching {ticker}..."):
        try:
            company = load_company(ticker)
        except Exception as e:  # network / rate limit
            st.error(f"Could not fetch data: {e}")
            st.stop()

    info = company["info"]
    d = derive(company)
    price = info.get("currentPrice") or info.get("regularMarketPrice") or info.get("previousClose")
    shares = info.get("sharesOutstanding")
    if not price or not shares:
        st.error("Price or share count missing for this ticker. Check the symbol / exchange suffix.")
        st.stop()

    model_choice = st.radio("Forecast model", ["Revenue & margin driven", "FCF growth (fallback)"],
                            help="Revenue & margin is the default. The fallback projects reported free "
                                 "cash flow directly, for companies with missing income-statement data.")
    fallback = model_choice.startswith("FCF")
    if not fallback and not d.has_operating:
        st.warning("Revenue / EBIT history unavailable, using the FCF-growth fallback.")
        fallback = True
    if fallback and d.fcf_hist.empty:
        st.error("No usable financial history for this ticker.")
        st.stop()

    st.header("2. Forecast")
    years = st.slider("Projection years", 3, 15, 10)
    tv_label = st.selectbox(
        "Terminal value method",
        ["Gordon growth", "Exit multiple (EV/EBITDA)", "Blend (average of both)"],
        disabled=fallback,
        help="The other method is always shown as a cross-check.")
    tv_method = "gordon" if fallback else {"Gordon": "gordon", "Exit": "exit", "Blend": "blend"}[tv_label.split()[0]]
    terminal_g = st.slider("Terminal growth (%)", 0.0, 6.0, 2.5, 0.1) / 100
    own_mult = info.get("enterpriseToEbitda")
    mult_default = float(np.clip(own_mult, 6, 25)) if isinstance(own_mult, (int, float)) and own_mult > 0 else 12.0
    exit_mult = st.number_input("Exit EV/EBITDA multiple (x)", 1.0, 60.0, round(mult_default, 1), 0.5,
                                key=K("mult"), disabled=fallback,
                                help="Defaults to the company's current multiple, capped 6-25x.")

    st.header("3. Discount rate")
    wacc_mode = st.radio("WACC", ["Calculate (CAPM)", "Manual"], horizontal=True)
    rf = st.number_input("Risk-free rate (%)", 0.0, 15.0, 4.2, 0.1) / 100
    erp = st.number_input("Equity risk premium (%)", 0.0, 12.0, 5.0, 0.1) / 100
    cost_debt = st.number_input("Pre-tax cost of debt (%)", 0.0, 20.0, 5.5, 0.1) / 100
    tax = st.number_input("Tax rate (%)", 0.0, 50.0, round(d.tax * 100, 1), 0.5, key=K("tax")) / 100
    manual_wacc = st.slider("Manual WACC (%)", 4.0, 20.0, 9.0, 0.1) / 100

    st.header("4. Decision rule")
    band = st.slider("Margin of safety band (±%)", 5, 40, 15,
                     help="Inside this band the stock is 'fairly valued'.") / 100

    st.header("5. Monte Carlo")
    n_sims = st.select_slider("Simulations", [1000, 2000, 5000, 10000, 20000], value=5000)
    seed = st.number_input("Random seed", 0, 99999, 42, 1)
    with st.expander("Input uncertainty (1 st. dev.)"):
        sd = {
            "growth": st.number_input("Revenue growth (pp)", 0.0, 15.0, 2.0, 0.5) / 100,
            "margin": st.number_input("Target margin (pp)", 0.0, 15.0, 1.5, 0.5) / 100,
            "wacc": st.number_input("WACC (pp)", 0.0, 4.0, 0.75, 0.25) / 100,
            "tg": st.number_input("Terminal growth (pp)", 0.0, 2.0, 0.4, 0.1) / 100,
            "mult_rel": st.number_input("Exit multiple (% of base)", 0.0, 50.0, 10.0, 1.0) / 100,
        }
    if fallback:
        sd["margin"] = 0.0
        sd["mult_rel"] = 0.0

    st.header("6. Peers")
    peers_txt = st.text_area("Peer tickers (comma separated)", key=K("peers"), height=80,
                             value=", ".join(DEFAULT_PEERS.get(ticker, [])),
                             help="Used for P/E and EV/EBITDA relative valuation.")

# ==========================================================================
# Company header
# ==========================================================================
currency = info.get("currency", "USD")
cur = {"USD": "$", "INR": "₹", "GBP": "£", "EUR": "€", "JPY": "¥"}.get(currency, currency + " ")
name = info.get("longName") or info.get("shortName") or ticker
beta = float(info.get("beta") or 1.0)
mcap = float(info.get("marketCap") or price * shares)
debt = float(info.get("totalDebt") or 0)
cash = float(info.get("totalCash") or 0)
net_debt = debt - cash
fin_cur = info.get("financialCurrency")

st.subheader(f"{name} ({ticker})")
if fin_cur and fin_cur != currency:
    st.warning(f"Financial statements are in {fin_cur} but the share price is in {currency}. "
               "Convert inputs to one currency before trusting the valuation.")
m1, m2, m3, m4 = st.columns(4)
m1.metric("Current price", f"{cur}{price:,.2f}")
m2.metric("Market cap", fmt_big(mcap, cur))
m3.metric("Beta", f"{beta:.2f}")
m4.metric("Net debt", fmt_big(net_debt, cur))

# ==========================================================================
# Operating assumptions
# ==========================================================================
with st.expander("Operating assumptions", expanded=True):
    if fallback:
        c1, c2 = st.columns(2)
        base_fcf = float(d.fcf_hist.iloc[-1])
        rev0 = c1.number_input(f"Base free cash flow ({cur}, FY{d.fcf_year})", value=base_fcf,
                               step=max(abs(base_fcf) * 0.01, 1.0), format="%.0f", key=K("fcf0"))
        hist_g = None
        pos = d.fcf_hist[d.fcf_hist > 0]
        if len(pos) >= 2 and pos.index[-1] > pos.index[0]:
            hist_g = (pos.iloc[-1] / pos.iloc[0]) ** (1 / (pos.index[-1] - pos.index[0])) - 1
        g_def = float(np.clip(hist_g if hist_g is not None else 0.08, -0.05, 0.30))
        growth = c2.slider("FCF growth, year 1 (%)", -10.0, 40.0, round(g_def * 100, 1), 0.5, key=K("g")) / 100
        if rev0 <= 0:
            st.warning("Base FCF is zero or negative; a DCF is unreliable. Override with a normalised figure.")
        a = Assumptions(rev0=rev0, growth_start=growth, margin_start=1.0, margin_target=1.0, margin_years=1,
                            tax=0.0, da_pct=0.0, capex_pct=0.0, nwc_pct=0.0, wacc=0.0, terminal_g=terminal_g,
                            exit_multiple=exit_mult, years=years, net_debt=net_debt, shares=float(shares),
                            tv_method="gordon")
    else:
        c1, c2, c3 = st.columns(3)
        rev0 = c1.number_input(f"Base revenue ({cur}, latest FY)", value=float(d.rev0),
                               step=max(abs(d.rev0) * 0.01, 1.0), format="%.0f", key=K("rev0"))
        growth = c1.slider("Revenue growth, year 1 (%)", -10.0, 40.0, round(d.growth * 100, 1), 0.5,
                           key=K("g"), help="Fades linearly to the terminal growth rate.") / 100
        tgt_def = float(np.clip(d.margin_target * 100, -10.0, 60.0))
        margin_target = c2.slider("Target EBIT margin (%)", -10.0, 60.0, round(tgt_def, 1), 0.5, key=K("mt"),
                                  help="Defaults to the 3-year average margin.") / 100
        c2.caption(f"Current EBIT margin: {d.margin_now:.1%}")
        margin_years = c2.slider("Years to reach target margin", 1, 15, min(5, years), key=K("my"))
        da_pct = c3.number_input("D&A (% of revenue)", 0.0, 30.0, round(d.da_pct * 100, 1), 0.1, key=K("da")) / 100
        capex_pct = c3.number_input("Capex (% of revenue)", 0.0, 40.0, round(d.capex_pct * 100, 1), 0.1,
                                    key=K("cx")) / 100
        nwc_pct = c3.number_input("NWC investment (% of revenue change)", 0.0, 50.0, round(d.nwc_pct * 100, 1),
                                  0.5, key=K("nwc")) / 100
        if capex_pct < da_pct:
            st.warning("Capex is below D&A. Final-year FCF (used for terminal value) may be overstated.")
        a = Assumptions(rev0=rev0, growth_start=growth, margin_start=d.margin_now,
                            margin_target=margin_target, margin_years=margin_years, tax=tax, da_pct=da_pct,
                            capex_pct=capex_pct, nwc_pct=nwc_pct, wacc=0.0, terminal_g=terminal_g,
                            exit_multiple=exit_mult, years=years, net_debt=net_debt, shares=float(shares),
                            tv_method=tv_method)

# ---- WACC ---------------------------------------------------------------
if wacc_mode.startswith("Calculate"):
    wacc, ke, kd, we, wd = compute_wacc(rf, beta, erp, cost_debt, tax, mcap, debt)
    with st.expander("WACC build-up"):
        st.write(f"Cost of equity (CAPM): **{ke:.2%}** = {rf:.2%} + {beta:.2f} × {erp:.2%}")
        st.write(f"After-tax cost of debt: **{kd:.2%}**; weights: equity **{we:.1%}**, debt **{wd:.1%}**")
        st.write(f"WACC: **{wacc:.2%}**")
else:
    wacc = manual_wacc
if wacc <= terminal_g:
    st.error("WACC must exceed terminal growth. Adjust the sidebar inputs.")
    st.stop()
a = replace(a, wacc=wacc)

# ==========================================================================
# Base case
# ==========================================================================
base = run_dcf(a)
iv = float(base["per_share"][0])
upside = iv / price - 1
if upside > band:
    verdict, vcolor, vicon, vhex = "UNDERVALUED", "green", "🟢", "#1E8449"
elif upside < -band:
    verdict, vcolor, vicon, vhex = "OVERVALUED", "red", "🔴", "#C0392B"
else:
    verdict, vcolor, vicon, vhex = "FAIRLY VALUED", "orange", "🟡", "#B9770E"

st.divider()
v1, v2, v3 = st.columns([2, 1, 1])
v1.markdown(f"### {vicon} Base-case verdict: :{vcolor}[{verdict}]")
v2.metric("Intrinsic value / share", f"{cur}{iv:,.2f}")
v3.metric("Upside / (downside)", f"{upside:+.1%}", delta=f"{cur}{iv - price:,.2f}")
if iv <= 0:
    st.warning("Intrinsic value is negative; cash flows or leverage make this DCF unreliable.")

tabs = st.tabs(["Valuation", "Forecast", "Scenarios", "Monte Carlo",
                "Relative valuation", "Sensitivity", "Export"])
t_val, t_fc, t_scn, t_mc, t_rel, t_sen, t_exp = tabs
last_fy = int(d.rev_hist.index[-1]) if (not fallback and len(d.rev_hist)) else int(d.fcf_year)
year_labels = [f"{last_fy + i}E" for i in range(1, years + 1)]

# ==========================================================================
# Forecast tab
# ==========================================================================
with t_fc:
    if fallback:
        fc = pd.DataFrame({
            "Year": year_labels, "Growth": base["growth"][0], "FCF": base["fcff"][0],
            "Discount factor": 1 / base["disc"][0], "PV of FCF": base["pv"][0]})
        fc_spec = {"Growth": "pct", "FCF": "big", "Discount factor": "num", "PV of FCF": "big"}
    else:
        fc = pd.DataFrame({
            "Year": year_labels, "Revenue": base["rev"][0], "Growth": base["growth"][0],
            "EBITDA": base["ebitda"][0], "EBIT": base["ebit"][0], "EBIT margin": base["margin"][0],
            "Capex": base["capex"][0], "NWC change": base["d_nwc"][0], "FCFF": base["fcff"][0],
            "PV of FCFF": base["pv"][0]})
        fc_spec = {"Revenue": "big", "Growth": "pct", "EBITDA": "big", "EBIT": "big", "EBIT margin": "pct",
                   "Capex": "big", "NWC change": "big", "FCFF": "big", "PV of FCFF": "big"}
    st.dataframe(format_df(fc, fc_spec, cur), hide_index=True)

    if not fallback:
        fig = go.Figure()
        hx = [str(y) for y in d.rev_hist.index]
        fig.add_bar(x=hx, y=d.rev_hist.values, name="Revenue (reported)", marker_color=GREY)
        fig.add_bar(x=year_labels, y=base["rev"][0], name="Revenue (forecast)", marker_color=BLUE)
        mh = d.margin_hist.reindex(d.rev_hist.index)
        fig.add_scatter(x=hx + year_labels, y=list(mh.values) + list(base["margin"][0]), name="EBIT margin",
                        yaxis="y2", mode="lines+markers", line=dict(color=RED))
        fig.update_layout(height=380, title="Revenue and EBIT margin", yaxis_title=f"Revenue ({currency})",
                          yaxis2=dict(title="EBIT margin", overlaying="y", side="right", tickformat=".0%"),
                          legend=dict(orientation="h", y=-0.2))
        st.plotly_chart(fig)

    fig2 = go.Figure()
    if len(d.fcf_hist):
        fig2.add_bar(x=[str(y) for y in d.fcf_hist.index], y=d.fcf_hist.values, name="FCF (reported)",
                     marker_color=GREY)
    fig2.add_bar(x=year_labels, y=base["fcff"][0], name="FCF (forecast)", marker_color=BLUE)
    fig2.update_layout(height=340, title="Free cash flow: reported vs forecast (sanity check)",
                       yaxis_title=f"FCF ({currency})", barmode="group", legend=dict(orientation="h", y=-0.2))
    st.plotly_chart(fig2)
    if not fallback:
        st.caption("FCFF = EBIT × (1 − tax) + D&A − capex − investment in NWC. If forecast FCF jumps far away "
                   "from reported FCF in year 1, revisit the margin, capex or NWC assumptions.")

# ==========================================================================
# Scenarios tab
# ==========================================================================
with t_scn:
    st.write("Edit the shifts applied to the base case. Probabilities are normalised to 100%.")
    C_G, C_M, C_W, C_T, C_X, C_P = ("Revenue growth Δ (pp)", "Target margin Δ (pp)", "WACC Δ (pp)",
                                    "Terminal g Δ (pp)", "Exit multiple Δ (x)", "Probability (%)")
    default_scn = pd.DataFrame({
        "Scenario": ["Bear", "Base", "Bull"], C_G: [-3.0, 0.0, 3.0], C_M: [-2.0, 0.0, 2.0],
        C_W: [1.0, 0.0, -1.0], C_T: [-0.5, 0.0, 0.5], C_X: [-2.0, 0.0, 2.0], C_P: [25.0, 50.0, 25.0]})
    cc = {c: st.column_config.NumberColumn(c, format="%.1f", step=0.5) for c in default_scn.columns[1:]}
    cc["Scenario"] = st.column_config.TextColumn("Scenario", disabled=True)
    scn_in = st.data_editor(default_scn, hide_index=True, column_config=cc, num_rows="fixed", key=K("scn"))
    if fallback:
        st.caption("Fallback mode: margin and exit-multiple shifts have no effect.")

    rows = []
    for _, r in scn_in.iterrows():
        sa = apply_scenario(a, r[C_G] / 100, 0.0 if fallback else r[C_M] / 100, r[C_W] / 100,
                                r[C_T] / 100, r[C_X])
        rs = run_dcf(sa)
        v = float(rs["per_share"][0])
        rows.append({"Scenario": r["Scenario"], "Probability": float(r[C_P]), "Revenue growth Y1": sa.growth_start,
                     "Target margin": sa.margin_target, "WACC": sa.wacc, "Terminal g": sa.terminal_g,
                     "Exit multiple": sa.exit_multiple, "Value / share": v, "Upside": v / price - 1})
    scn = pd.DataFrame(rows)
    pw = scn["Probability"].where(scn["Value / share"].notna(), 0.0)
    weighted = float((scn["Value / share"].fillna(0) * pw).sum() / pw.sum()) if pw.sum() > 0 else np.nan
    scn["Probability"] = scn["Probability"] / scn["Probability"].sum()
    scn_spec = {"Probability": "pct", "Revenue growth Y1": "pct", "Target margin": "pct", "WACC": "pct",
                "Terminal g": "pct", "Exit multiple": "x", "Value / share": "price", "Upside": "pct"}
    shown = scn.drop(columns=["Target margin", "Exit multiple"]) if fallback else scn
    st.dataframe(format_df(shown, scn_spec, cur), hide_index=True)

    s1, s2 = st.columns(2)
    s1.metric("Probability-weighted value", f"{cur}{weighted:,.2f}" if np.isfinite(weighted) else "n/a",
              delta=f"{weighted / price - 1:+.1%} vs price" if np.isfinite(weighted) else None)
    s2.metric("Current price", f"{cur}{price:,.2f}")
    figs = go.Figure(go.Bar(x=scn["Scenario"], y=scn["Value / share"],
                            marker_color=[RED if (x or 0) < price else BLUE for x in scn["Value / share"]],
                            text=[f"{cur}{x:,.2f}" if np.isfinite(x) else "n/a" for x in scn["Value / share"]],
                            textposition="outside"))
    figs.add_hline(y=price, line_dash="dash", line_color=NAVY, annotation_text=f"Price {cur}{price:,.2f}")
    figs.update_layout(height=360, yaxis_title=f"Value per share ({currency})", showlegend=False)
    st.plotly_chart(figs)

scn_vals = scn["Value / share"].dropna()

# ==========================================================================
# Monte Carlo tab
# ==========================================================================
with t_mc:
    samples = monte_carlo(a, int(n_sims), int(seed), sd)
    if samples.size < 100:
        st.error("Too few valid simulations; widen WACC vs terminal growth headroom.")
        mc = None
    else:
        mc = mc_summary(samples, price, band)
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Median value / share", f"{cur}{mc['pct'][50]:,.2f}")
        k2.metric("Mean value / share", f"{cur}{mc['mean']:,.2f}")
        k3.metric("P(value > price)", f"{mc['p_above']:.0%}")
        k4.metric(f"P(upside > {band:.0%})", f"{mc['p_above_band']:.0%}")

        lo, hi = np.percentile(samples, [1, 99])
        figm = go.Figure(go.Histogram(x=np.clip(samples, lo, hi), nbinsx=60, marker_color=BLUE))
        figm.add_vline(x=price, line_dash="dash", line_color=RED, annotation_text=f"Price {cur}{price:,.2f}")
        figm.add_vline(x=mc["pct"][50], line_color=NAVY, annotation_text="Median", annotation_position="top left")
        figm.update_layout(height=380, xaxis_title=f"Intrinsic value per share ({currency}), clipped at P1-P99",
                           yaxis_title="Simulations", showlegend=False)
        st.plotly_chart(figm)

        mc_tbl = pd.DataFrame({"Statistic": ["P5", "P10", "P25", "Median", "P75", "P90", "P95", "Mean", "Std. dev."],
                               "Value / share": [mc["pct"][p] for p in (5, 10, 25, 50, 75, 90, 95)]
                               + [mc["mean"], mc["std"]]})
        mc_spec = {"Value / share": "price"}
        st.dataframe(format_df(mc_tbl, mc_spec, cur), hide_index=True)
        st.caption(f"{mc['n']:,} valid simulations. Inputs are drawn independently from normal distributions "
                   "around the base case and clipped to sane bounds; real-world inputs are correlated "
                   "(for example growth and margins), so treat the spread as indicative.")

# ==========================================================================
# Relative valuation tab
# ==========================================================================
peer_list = [p for p in dict.fromkeys(re.split(r"[,\s]+", peers_txt.upper())) if p and p != ticker][:10]
rel_rows, rel_tbl, peers_df = [], None, None
with t_rel:
    if not peer_list:
        st.info("Add peer tickers in the sidebar (section 6) to value the company on peer P/E and EV/EBITDA.")
    else:
        with st.spinner("Fetching peer multiples..."):
            peers_df = load_multiples(tuple(peer_list))
        own = pd.DataFrame([{"Ticker": ticker, "Name": name[:28], "Market cap": mcap,
                             "P/E (TTM)": info.get("trailingPE"), "P/E (Fwd)": info.get("forwardPE"),
                             "EV/EBITDA": info.get("enterpriseToEbitda"), "EV/Revenue": info.get("enterpriseToRevenue")}])
        for c in own.columns[2:]:
            own[c] = pd.to_numeric(own[c], errors="coerce")
        peer_spec = {"Market cap": "big", "P/E (TTM)": "x", "P/E (Fwd)": "x", "EV/EBITDA": "x", "EV/Revenue": "x"}
        st.dataframe(format_df(pd.concat([own, peers_df], ignore_index=True), peer_spec, cur), hide_index=True)
        st.caption("First row is the company itself. Negative or missing multiples are ignored in the statistics.")

        def q(col, p):
            s = peers_df[col].replace([np.inf, -np.inf], np.nan)
            s = s[s > 0].dropna()
            return float(np.percentile(s, p)) if len(s) else np.nan

        eps = d.eps
        ebitda = d.ebitda_ttm
        for label, col, kind in (("P/E (TTM)", "P/E (TTM)", "pe"), ("EV/EBITDA", "EV/EBITDA", "ev")):
            lo_m, md_m, hi_m = q(col, 25), q(col, 50), q(col, 75)
            vals = [np.nan] * 3
            if kind == "pe" and np.isfinite(eps) and eps > 0:
                vals = [m * eps for m in (lo_m, md_m, hi_m)]
            if kind == "ev" and np.isfinite(ebitda) and ebitda > 0:
                vals = [(m * ebitda - net_debt) / shares for m in (lo_m, md_m, hi_m)]
            rel_rows.append({"Method": label, "Peer P25": lo_m, "Peer median": md_m, "Peer P75": hi_m,
                             "Value @ P25": vals[0], "Value @ median": vals[1], "Value @ P75": vals[2]})
        rel_tbl = pd.DataFrame(rel_rows)
        rel_spec = {"Peer P25": "x", "Peer median": "x", "Peer P75": "x",
                    "Value @ P25": "price", "Value @ median": "price", "Value @ P75": "price"}
        st.markdown("**Implied value per share from peer multiples**")
        st.dataframe(format_df(rel_tbl, rel_spec, cur), hide_index=True)
        st.caption(f"Company TTM EPS: {eps:,.2f}" if np.isfinite(eps) else "Company EPS unavailable (or negative).")
        if not (np.isfinite(ebitda) and ebitda > 0):
            st.caption("Company EBITDA unavailable or negative, so EV/EBITDA value is not computed.")

# ==========================================================================
# Sensitivity tab
# ==========================================================================
with t_sen:
    axis_label = st.radio("Column axis", ["Terminal growth", "Exit multiple"], horizontal=True,
                          index=1 if tv_method == "exit" else 0, disabled=fallback)
    axis = "tg" if (fallback or axis_label == "Terminal growth") else "mult"
    waccs, cols, grid = sensitivity(a, axis)
    col_names = [f"{c:.1%}" if axis == "tg" else f"{c:.1f}x" for c in cols]
    sens = pd.DataFrame(grid, columns=col_names)
    sens.insert(0, "WACC", waccs)
    sens_spec = {"WACC": "pct", **{c: "price" for c in col_names}}
    if tv_method == "gordon" and axis == "mult":
        st.info("Terminal value uses Gordon growth, so the exit multiple has no effect. Switch the method "
                "or the axis.")
    if tv_method == "exit" and axis == "tg":
        st.info("Terminal value uses the exit multiple, so terminal growth has no effect. Switch the axis.")

    def shade(v):
        if pd.isna(v):
            return ""
        return ("background-color:#d4edda;color:#155724" if v > price else "background-color:#f8d7da;color:#721c24")

    styled = (sens.style.format({c: f"{cur}{{:,.2f}}" for c in col_names} | {"WACC": "{:.1%}"})
              .map(shade, subset=col_names).hide(axis="index"))
    st.dataframe(styled)
    st.caption(f"Value per share. Green = above the current price ({cur}{price:,.2f}), red = below.")

# ==========================================================================
# Valuation tab (needs everything above)
# ==========================================================================
football = []
if len(scn_vals):
    football.append(("DCF: bear to bull", float(scn_vals.min()), iv, float(scn_vals.max())))
if mc is not None:
    football.append(("Monte Carlo P10-P90", mc["pct"][10], mc["pct"][50], mc["pct"][90]))
for r in rel_rows:
    if np.isfinite(r["Value @ P25"]) and np.isfinite(r["Value @ P75"]):
        football.append((f"Peers {r['Method']} (P25-P75)", r["Value @ P25"], r["Value @ median"], r["Value @ P75"]))
lo52, hi52 = info.get("fiftyTwoWeekLow"), info.get("fiftyTwoWeekHigh")
if lo52 and hi52:
    football.append(("52-week range", float(lo52), price, float(hi52)))

with t_val:
    bridge = pd.DataFrame({
        "Component": ["PV of forecast cash flows", "PV of terminal value", "Enterprise value", "Less: net debt",
                      "Equity value", "Shares outstanding", "Intrinsic value / share"],
        "Value": [base["pv_sum"][0], base["pv_tv"][0], base["ev"][0], -net_debt, base["equity"][0],
                  float(shares), iv]})
    brs = format_df(bridge, {"Value": "big"}, cur)
    brs.loc[brs["Component"] == "Intrinsic value / share", "Value"] = f"{cur}{iv:,.2f}"
    st.markdown("**Enterprise to equity bridge**")
    st.dataframe(brs, hide_index=True)
    st.caption(f"Terminal value is {base['tv_share'][0]:.0%} of enterprise value; a high share means the "
               "result leans heavily on long-run assumptions.")

    xc = None
    if fallback:
        st.info("Exit-multiple cross-check needs EBITDA, so it is available in the revenue & margin model.")
    else:
        rg = run_dcf(replace(a, tv_method="gordon"))
        rx = run_dcf(replace(a, tv_method="exit"))
        vg, vx = float(rg["per_share"][0]), float(rx["per_share"][0])
        xc = pd.DataFrame({
            "Terminal value method": ["Gordon growth", "Exit multiple (EV/EBITDA)"],
            "Value / share": [vg, vx], "vs price": [vg / price - 1, vx / price - 1],
            "TV % of EV": [float(rg["tv_share"][0]), float(rx["tv_share"][0])],
            "Terminal growth": [a.terminal_g, float(rx["implied_g"][0])],
            "Exit EV/EBITDA": [float(rg["implied_mult"][0]), a.exit_multiple]})
        xc_spec = {"Value / share": "price", "vs price": "pct", "TV % of EV": "pct",
                   "Terminal growth": "pct", "Exit EV/EBITDA": "x"}
        st.markdown("**Terminal value cross-check**")
        st.dataframe(format_df(xc, xc_spec, cur), hide_index=True)
        st.caption("Gordon row shows the exit multiple its terminal value implies; the exit-multiple row shows "
                   "the perpetual growth rate its terminal value implies.")
        gap = abs(vx / vg - 1) if vg else np.nan
        ig = float(rx["implied_g"][0])
        if np.isfinite(ig) and ig > min(wacc - 0.01, 0.05):
            st.warning(f"The exit multiple implies perpetual growth of {ig:.1%}, which is aggressive. "
                       "Consider a lower multiple.")
        elif np.isfinite(gap) and gap > 0.25:
            st.warning(f"The two terminal value methods differ by {gap:.0%}. Revisit the exit multiple or "
                       "terminal growth until they are reasonably consistent.")

    if football:
        st.markdown("**Valuation summary (football field)**")
        figf = go.Figure()
        for label, lo, mid, hi in football:
            figf.add_bar(y=[label], x=[hi - lo], base=[lo], orientation="h", marker_color=BLUE, showlegend=False,
                         hovertemplate=f"{label}<br>{cur}{lo:,.2f} to {cur}{hi:,.2f}<extra></extra>")
            figf.add_scatter(x=[mid], y=[label], mode="markers", showlegend=False,
                             marker=dict(symbol="line-ns", size=22, line=dict(width=3, color="white")),
                             hovertemplate=f"Mid {cur}{mid:,.2f}<extra></extra>")
        figf.add_vline(x=price, line_dash="dash", line_color=RED, annotation_text=f"Price {cur}{price:,.2f}")
        figf.update_layout(height=120 + 55 * len(football), xaxis_title=f"Value per share ({currency})",
                           yaxis=dict(autorange="reversed"), bargap=0.45)
        st.plotly_chart(figf)

# ==========================================================================
# Export tab
# ==========================================================================
with t_exp:
    sig = (ticker, round(iv, 4), int(n_sims), int(seed), fallback, tv_method, round(wacc, 5))
    if st.button("Generate Excel and PDF report", type="primary"):
        assump_rows = [("Forecast model", "FCF growth (fallback)" if fallback else "Revenue & margin"),
                       ("Projection years", years), ("Base " + ("FCF" if fallback else "revenue"), fmt_big(a.rev0, cur)),
                       ("Growth, year 1", f"{a.growth_start:.1%}"), ("Terminal growth", f"{a.terminal_g:.1%}"),
                       ("WACC", f"{wacc:.2%}"), ("Tax rate", f"{tax:.1%}")]
        if not fallback:
            assump_rows += [("Current EBIT margin", f"{a.margin_start:.1%}"),
                            ("Target EBIT margin", f"{a.margin_target:.1%} in {a.margin_years} yrs"),
                            ("D&A % revenue", f"{a.da_pct:.1%}"), ("Capex % revenue", f"{a.capex_pct:.1%}"),
                            ("NWC % of revenue change", f"{a.nwc_pct:.1%}"),
                            ("Terminal value method", tv_label), ("Exit EV/EBITDA", f"{a.exit_multiple:.1f}x")]
        assump_rows += [("Net debt", fmt_big(net_debt, cur)), ("Shares outstanding", fmt_big(float(shares))),
                        ("Beta", f"{beta:.2f}"), ("Margin-of-safety band", f"±{band:.0%}")]

        summary = [("Company", f"{name} ({ticker})"), ("Current price", f"{cur}{price:,.2f}"),
                   ("Base-case intrinsic value / share", f"{cur}{iv:,.2f}"), ("Upside / (downside)", f"{upside:+.1%}"),
                   ("Verdict", verdict)]
        if np.isfinite(weighted):
            summary.append(("Probability-weighted scenario value", f"{cur}{weighted:,.2f}"))
        if len(scn_vals):
            summary.append(("Bear / Bull value", f"{cur}{scn_vals.min():,.2f} / {cur}{scn_vals.max():,.2f}"))
        if mc is not None:
            summary += [("Monte Carlo median", f"{cur}{mc['pct'][50]:,.2f}"),
                        ("Monte Carlo P(value > price)", f"{mc['p_above']:.0%}")]
        for r in rel_rows:
            if np.isfinite(r["Value @ median"]):
                summary.append((f"Peer {r['Method']} implied value", f"{cur}{r['Value @ median']:,.2f}"))

        tables = [
            {"name": "Assumptions", "df": pd.DataFrame(assump_rows, columns=["Parameter", "Value"]),
             "spec": {}, "pdf": True},
            {"name": "Projections", "df": fc, "spec": fc_spec, "pdf": True},
            {"name": "Scenarios", "df": scn, "spec": scn_spec, "pdf": True},
        ]
        if xc is not None:
            tables.append({"name": "TV cross-check", "df": xc, "spec": xc_spec, "pdf": True})
        tables.append({"name": "Sensitivity", "df": sens, "spec": sens_spec, "pdf": True})
        if mc is not None:
            tables.append({"name": "Monte Carlo", "df": mc_tbl, "spec": mc_spec, "pdf": True})
        if rel_tbl is not None:
            tables.append({"name": "Relative valuation", "df": rel_tbl, "spec": rel_spec, "pdf": True})
            tables.append({"name": "Peer multiples", "df": peers_df, "spec": peer_spec, "pdf": True})
        if mc is not None:
            tables.append({"name": "MC samples", "df": pd.DataFrame({"Value / share": samples[:5000]}),
                           "spec": {"Value / share": "price"}, "pdf": False})

        report = {"title": "DCF Equity Valuation Report", "ticker": ticker, "name": name, "currency": currency,
                  "date": date.today().isoformat(), "price": float(price), "verdict": verdict,
                  "verdict_color": vhex, "summary": summary, "tables": tables, "football": football,
                  "mc_samples": samples if mc is not None else None}
        with st.spinner("Building files..."):
            st.session_state["report_files"] = {
                "sig": sig, "xlsx": build_excel(report), "pdf": build_pdf(report)}

    files = st.session_state.get("report_files")
    if files is None:
        st.info("Click the button to build the files for the current inputs.")
    else:
        if files["sig"] != sig:
            st.warning("Inputs changed since these files were built. Regenerate to refresh them.")
        e1, e2 = st.columns(2)
        e1.download_button("⬇️ Download Excel (.xlsx)", files["xlsx"], file_name=f"{ticker}_dcf_report.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        e2.download_button("⬇️ Download PDF", files["pdf"], file_name=f"{ticker}_dcf_report.pdf",
                           mime="application/pdf")
        st.caption("The Excel file holds values (not live formulas), with one sheet per table.")

st.divider()
st.caption("Data from Yahoo Finance via yfinance; it may be delayed or incomplete. DCF output is only as good "
           "as its assumptions. Use it alongside qualitative research. Not investment advice.")
