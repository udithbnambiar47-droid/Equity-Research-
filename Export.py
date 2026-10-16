"""
Formatting and report export (Excel and PDF).

A report is a plain dict assembled by app.py:

    {
      "title", "ticker", "name", "currency", "date",
      "price", "verdict", "summary": [(label, value_str), ...],
      "tables": [{"name", "df", "spec": {col: kind}, "pdf": bool}, ...],
      "football": [(label, low, mid, high), ...],
      "mc_samples": np.ndarray | None,
    }

`spec` kinds: big | price | num | pct | x | int | txt
"""

from __future__ import annotations

import io
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

import engine as eng  # noqa: E402

XL_FORMATS = {"big": "#,##0", "price": "#,##0.00", "num": "#,##0.00",
              "pct": "0.0%", "x": '0.0"x"', "int": "0", "txt": "@"}


def _fmt_val(v, kind: str, cur: str = "") -> str:
    if v is None or (isinstance(v, (float, np.floating)) and not np.isfinite(v)):
        return "n/a"
    if pd.isna(v) if not isinstance(v, str) else False:
        return "n/a"
    try:
        if kind == "big":
            return eng.fmt_big(float(v), cur)
        if kind == "price":
            return f"{cur}{float(v):,.2f}"
        if kind == "num":
            return f"{float(v):,.2f}"
        if kind == "pct":
            return f"{float(v):.1%}"
        if kind == "x":
            return f"{float(v):.1f}x"
        if kind == "int":
            return f"{int(v)}"
    except (TypeError, ValueError):
        pass
    return str(v)


def format_df(df: pd.DataFrame, spec: dict, cur: str = "") -> pd.DataFrame:
    """Return a copy of df as display strings according to spec."""
    out = pd.DataFrame(index=df.index)
    for c in df.columns:
        kind = spec.get(c, "txt")
        out[c] = [_fmt_val(v, kind, cur) for v in df[c]]
    return out


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def build_excel(report: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")

    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        summary = pd.DataFrame(report["summary"], columns=["Item", "Value"])
        summary.to_excel(xw, sheet_name="Summary", index=False)
        used = {"Summary"}
        specs = {"Summary": {}}

        for t in report["tables"]:
            name = t["name"][:31]
            while name in used:
                name = name[:28] + "_2"
            used.add(name)
            specs[name] = t["spec"]
            t["df"].to_excel(xw, sheet_name=name, index=False)

        for name, ws in xw.sheets.items():
            spec = specs.get(name, {})
            for cell in ws[1]:
                cell.fill, cell.font = head_fill, head_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            headers = [c.value for c in ws[1]]
            for j, h in enumerate(headers, start=1):
                kind = spec.get(h, "txt")
                width = len(str(h))
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell = row[0]
                    if kind != "txt" and isinstance(cell.value, (int, float)):
                        cell.number_format = XL_FORMATS[kind]
                    width = max(width, len(str(cell.value)) if cell.value is not None else 0)
                ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = min(max(width + 3, 10), 60)
    return bio.getvalue()


# --------------------------------------------------------------------------
# Charts for the PDF
# --------------------------------------------------------------------------
def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def football_png(rows, price, cur):
    fig, ax = plt.subplots(figsize=(7.2, 0.55 * len(rows) + 1.0))
    for i, (label, lo, mid, hi) in enumerate(rows):
        ax.barh(i, hi - lo, left=lo, color="#3B6FB6", height=0.5)
        if mid is not None and np.isfinite(mid):
            ax.plot([mid, mid], [i - 0.25, i + 0.25], color="white", lw=2)
        ax.text(hi, i, f"  {lo:,.0f} - {hi:,.0f}", va="center", fontsize=7)
    ax.axvline(price, color="#C0392B", ls="--", lw=1.2)
    ax.text(price, -0.8, f"Price {price:,.2f}", color="#C0392B", fontsize=7, ha="center")
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels([r[0] for r in rows], fontsize=8)
    ax.invert_yaxis()
    ax.set_xlabel(f"Value per share ({cur.strip()})", fontsize=8)
    ax.tick_params(axis="x", labelsize=7)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _png(fig)


def mc_png(samples, price, cur):
    lo, hi = np.percentile(samples, [1, 99])
    fig, ax = plt.subplots(figsize=(7.2, 2.6))
    ax.hist(np.clip(samples, lo, hi), bins=60, color="#3B6FB6")
    ax.axvline(price, color="#C0392B", ls="--", lw=1.2, label=f"Price {price:,.2f}")
    ax.axvline(np.median(samples), color="black", lw=1.2, label=f"Median {np.median(samples):,.2f}")
    ax.legend(fontsize=7, frameon=False)
    ax.set_xlabel(f"Intrinsic value per share ({cur.strip()})", fontsize=8)
    ax.tick_params(labelsize=7)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _png(fig)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def build_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    cur = report["currency"] + " "          # Helvetica has no currency glyphs, use the code
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=17, spaceAfter=2, alignment=0)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceBefore=10, spaceAfter=4,
                        textColor=colors.HexColor("#1F3A5F"))
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7.5, textColor=colors.grey)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9)

    def safe(x):
        """Helvetica only covers Latin-1; swap currency symbols and unknown glyphs."""
        x = str(x).replace("₹", "INR ").replace("€", "EUR ").replace("Δ", "Chg ")
        return x.encode("latin-1", "replace").decode("latin-1")

    def table(df, spec, col_align_left=1):
        fdf = format_df(df, spec, cur)
        data = [[safe(c) for c in fdf.columns]] + [[safe(v) for v in row] for row in fdf.values.tolist()]
        tb = Table(data, repeatRows=1, hAlign="LEFT")
        st_ = [
            ("FONT", (0, 0), (-1, -1), "Helvetica", 7.5),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (col_align_left, 0), (-1, -1), "RIGHT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
            ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.grey),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]
        tb.setStyle(TableStyle(st_))
        return tb

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=report["title"], author="DCF Equity Research App")
    story = [
        Paragraph(escape(safe(report["title"])), h1),
        Paragraph(escape(safe(f"{report['name']} ({report['ticker']})  |  {report['date']}  |  prices in {report['currency']}")), small),
        Spacer(1, 6),
        Paragraph(f"<b>Verdict: {escape(report['verdict'])}</b>", ParagraphStyle(
            "v", parent=body, fontSize=13, textColor=colors.HexColor(report.get("verdict_color", "#1F3A5F")))),
        Spacer(1, 4),
    ]

    sm = pd.DataFrame(report["summary"], columns=["Item", "Value"])
    story.append(table(sm, {}))

    if report.get("football"):
        story += [Paragraph("Valuation summary (football field)", h2),
                  Image(football_png(report["football"], report["price"], cur), width=175 * mm,
                        height=(0.55 * len(report["football"]) + 1.0) / 7.2 * 175 * mm * 0.97)]
    samples = report.get("mc_samples")
    if samples is not None and len(samples):
        story += [Paragraph("Monte Carlo distribution", h2),
                  Image(mc_png(samples, report["price"], cur), width=175 * mm, height=175 * mm * 2.6 / 7.2)]

    for t in report["tables"]:
        if not t.get("pdf", True):
            continue
        story.append(KeepTogether([Paragraph(escape(safe(t["name"])), h2), table(t["df"], t["spec"])]))

    story += [Spacer(1, 10), Paragraph(
        "Educational use only; not investment advice. Data from Yahoo Finance (yfinance) may be delayed "
        "or incomplete. Model output depends entirely on the stated assumptions.", small)]
    doc.build(story)
    return buf.getvalue()
